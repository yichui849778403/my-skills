# -*- coding: utf-8 -*-
"""结构校验：确认"去备注"版除了备注以外没被动过。

用法:
    python verify_no_notes.py <原文件.pptx> <去备注后.pptx>

校验项:
    1. zip 完整性
    2. 除 notesSlides/{rels,Content_Types,app.xml} 外，其余条目逐字节一致
    3. ppt/slides/_rels/*.rels 的差异恰好是各少 1 条 notesSlide 关系
    4. media 条目数与内容 MD5 完全一致
    5. 幻灯片数、正文文本框文本完全一致
    6. python-pptx 能正常打开，且 0 页有 notesSlide

注意 PowerPoint 侧的判读: 删掉备注后 Slides.Item(n).NotesPage 依然存在
（备注页是视图，由 notesMaster 派生），其中会残留"幻灯片编号"占位符
（PlaceholderFormat.Type = 13），显示为 '1'..'N'。这不算备注文本，
所以 COM 校验不能期望"有文本的备注页数 = 0"，要比的是字符总量急剧下降
（示例: 17075 -> 161，161 正好是 1..85 的编号长度）。
"""
import hashlib
import re
import sys
import zipfile

from pptx import Presentation

NOTES_SLIDE_TYPE = "notesSlide"


def verify(src, out):
    zi, zo = zipfile.ZipFile(src), zipfile.ZipFile(out)
    ok = True

    bad = zo.testzip()
    print("1) zip 完整性:", bad or "OK")
    ok &= bad is None

    diff = []
    checked = 0
    for n in zi.namelist():
        if n.startswith("ppt/notesSlides/"):
            continue
        checked += 1
        if zi.read(n) != zo.read(n):
            diff.append(n)
    only_expected = all(
        d in ("[Content_Types].xml", "docProps/app.xml")
        or d.startswith("ppt/slides/_rels/")
        for d in diff
    )
    print("2) 逐字节比对 %d 个条目，差异项全部为预期的备注相关文件: %s"
          % (checked, only_expected))
    ok &= only_expected

    rel_bad = []
    for n in zi.namelist():
        if not n.startswith("ppt/slides/_rels/"):
            continue
        ra = re.findall(r"<Relationship[^>]*/>", zi.read(n).decode("utf-8"))
        rb = re.findall(r"<Relationship[^>]*/>", zo.read(n).decode("utf-8"))
        only_a = [x for x in ra if x not in rb]
        only_b = [x for x in rb if x not in ra]
        if len(only_a) != 1 or NOTES_SLIDE_TYPE not in only_a[0] or only_b:
            rel_bad.append(n)
    print("3) rels 差异异常项:", rel_bad or "无（每页只少 1 条 notesSlide 关系）")
    ok &= not rel_bad

    media = lambda z: {
        n: hashlib.md5(z.read(n)).hexdigest() for n in z.namelist() if "/media/" in n
    }
    same_media = media(zi) == media(zo)
    print("4) media 条目 %d -> %d，内容一致: %s"
          % (len(media(zi)), len(media(zo)), same_media))
    ok &= same_media

    pa, pb = Presentation(src), Presentation(out)
    ta = [sh.text_frame.text for s in pa.slides for sh in s.shapes if sh.has_text_frame]
    tb = [sh.text_frame.text for s in pb.slides for sh in s.shapes if sh.has_text_frame]
    print("5) 幻灯片 %d -> %d，有备注页的页数 %d -> %d，正文文本一致: %s"
          % (len(pa.slides), len(pb.slides),
             sum(1 for s in pa.slides if s.has_notes_slide),
             sum(1 for s in pb.slides if s.has_notes_slide),
             ta == tb))
    ok &= ta == tb and len(pa.slides) == len(pb.slides)
    ok &= sum(1 for s in pb.slides if s.has_notes_slide) == 0

    print("RESULT:", "PASS" if ok else "FAIL")
    return ok


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    sys.exit(0 if verify(sys.argv[1], sys.argv[2]) else 2)
