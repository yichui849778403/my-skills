# -*- coding: utf-8 -*-
"""
数据分类分级 Excel 纠错重跑脚本（通用版）

背景：电信分类分级自动化工具导出的「分类分级清单」常过度分类（把 delete_flag、
      create_time 等字段误判为"个人信息/医疗健康信息"），原因是工具拿英文字段名
      硬猜，没用上它自己抽到的中文字段描述。

本脚本改用「中文字段描述 + 表描述」做关键词规则匹配，重跑分类分级，产出：
  1) <前缀>分类分级清单-修正版.xlsx   —— 全量字段，敏感行红/黄高亮
  2) <前缀>敏感字段清单-待人工确认.xlsx —— 仅三级+四级，供人工逐条拍板

用法：
  python classify_data.py [输入.xlsx] [输出前缀]
  默认输入 ./数据资产清单.xlsx，输出前缀默认空（即"分类分级清单-修正版.xlsx"）

依赖：openpyxl（pip install openpyxl）
列定位：按表头名自动匹配，列顺序变了也能跑；缺关键列会报错提示。

分级框架（四级）：
  一级不敏感（时间戳/id/标志位/统计数）/ 二级低敏感（一般业务数据）
  / 三级敏感（手机号/姓名/车牌等一般个人信息）/ 四级高敏感（身份证/密码等强身份鉴权）。

⚠️ 重要：下方敏感词词典只是通用兜底。不同单位系统字段差异大（医疗/教育/车管/交通…），
  不可能一套词覆盖所有单位。跑脚本前必须：先粗跑一遍 → 人工看漏判/误判 →
  把该单位的业务敏感词补进词典 → 重跑。人工过一遍敏感字段候选这步不能省（规则匹配必有边缘误判）。
"""

import sys
import re
import openpyxl
from collections import Counter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ============================================================
# 敏感词词典（核心，按单位定制 —— 每跑新单位先扩充这里）
# 高敏感(四级)：强身份/鉴权信息
# ============================================================
HIGH_CN = {
    '身份证': '个人身份信息', '驾驶证': '个人身份信息', '护照': '个人身份信息',
    '社保卡': '个人身份信息', '军官证': '个人身份信息', '行驶证': '个人身份信息',
    '银行卡': '个人财产信息', '银行账号': '个人财产信息',
    '证件号': '个人身份信息', '证件号码': '个人身份信息',
    '密码': '身份鉴别信息', '口令': '身份鉴别信息',
}
# 敏感(三级)：一般个人信息
SENS_CN = {
    '手机号': '个人通信信息', '手机号码': '个人通信信息', '联系电话': '个人通信信息',
    '手机': '个人通信信息', '电话': '个人通信信息', '联系方式': '个人通信信息',
    '传真': '个人通信信息', '邮箱': '个人通信信息',
    '姓名': '个人基本资料', '联系人': '个人基本资料', '性别': '个人基本资料',
    '出生': '个人基本资料', '生日': '个人基本资料',
    '住址': '个人位置信息', '家庭住址': '个人位置信息', '户籍': '个人位置信息',
    '车牌': '车辆标识信息', '号牌': '车辆标识信息',
    '车架号': '车辆标识信息', '发动机号': '车辆标识信息', '车辆识别': '车辆标识信息',
    '指纹': '个人生物识别信息', '人脸': '个人生物识别信息',
    '体检': '个人健康生理信息', '病历': '个人健康生理信息',
    '工资': '个人财产信息', '薪酬': '个人财产信息',
    'openid': '个人身份信息',
}
# 英文字段名关键词（兜底中文描述为空的字段）
# 注意：勿加 vin/engine 等短子串词，会误命中 province(pro"vin"ce)、driving(dri"vin"g)、engine_speed
HIGH_EN = ['id_card', 'idcard', 'sfz', 'password', 'passwd', 'pwd',
           'bankcard', 'driver_license', 'dlicense',
           'certifino', 'certno', 'id_no', 'idnum']
SENS_EN = ['phone', 'mobile', 'phon', 'email', 'plate', 'carno',
           'contactnumber', 'contact', 'birthday', 'openid']

# 建议管控措施映射
MEASURE = {
    '个人身份信息': '脱敏展示 + 加密存储 + 最小化访问 + 操作审计',
    '身份鉴别信息': '加密存储(hash/不可逆) + 禁止明文/日志 + 强口令策略',
    '个人通信信息': '脱敏展示 + 最小化访问 + 导出管控',
    '个人基本资料': '访问控制 + 脱敏(按需) + 操作审计',
    '车辆标识信息': '访问控制 + 操作审计 + 按需脱敏',
    '个人位置信息': '访问控制 + 最小化采集 + 按需脱敏',
    '个人财产信息': '加密存储 + 最小化访问 + 操作审计',
    '个人生物识别信息': '严格访问控制 + 加密存储 + 留存期限管控',
    '个人健康生理信息': '严格访问控制 + 加密存储 + 最小化',
    '非敏感数据': '无需特殊防护',
    '业务数据': '常规防护(身份认证/访问控制/备份/审计)',
}

# ============================================================
# 列定位（按表头名）
# ============================================================
REQUIRED = ['数据表名', '字段名称', '字段描述', '表描述']
OPTIONAL = ['ip', 'port', '实例名', '数据库名', '模式名']


def locate_columns(header):
    """按表头名定位列，返回 {列名: 索引(0起)}。缺关键列抛异常。"""
    idx = {name: i for i, name in enumerate(header) if name is not None}
    for col in REQUIRED:
        if col not in idx:
            raise KeyError(f'缺少关键列「{col}」，实际表头: {header}')
    return idx


def is_url_field(fname, desc):
    """接口/附件/发票等 URL 类字段不算个人地址。"""
    text = (fname or '').lower() + (desc or '')
    for kw in ('url', '接口', '网关', '附件', '发票', '保单', '申请函', '证明',
               '图片', '照片', '文件'):
        if kw in text:
            return True
    return False


def is_level1(fname, desc):
    """判断一级（时间戳/id/标志位/统计数，无个人信息、可公开）。返回 (bool, 依据)。"""
    f = (fname or '').lower().strip()
    d = (desc or '')
    # 主键/外键 id
    if f == 'id':
        return True, '主键id'
    if f.endswith('_id'):
        return (True, '外键id') if (not d or 'id' in d.lower()) else (False, '')
    # 时间戳（中文注释）
    for kw in ('创建时间', '修改时间', '更新时间', '时间戳', '创建日期', '修改日期',
               '注册时间', '最后登录', '绑定时间', '登记日期', '接收时间', '删除日期',
               '领证日期', '换证时间', '年检日期', '保养日期', '保险日期', '购置日期'):
        if kw in d:
            return True, f'命中[{kw}]'
    # 时间戳（英文字段名兜底）
    if f.endswith('_time') or f.endswith('_date') or f.endswith('_at'):
        return True, '字段名时间模式'
    if f in ('timestamp', 'createtime', 'updatetime', 'createdon', 'modifedon',
             'modifiedon', 'createdat', 'updatedat'):
        return True, '字段名时间模式'
    # 标志位（中文注释）
    for kw in ('是否', '标记', '标志', '删除标识', '有效标记'):
        if kw in d:
            return True, f'命中[{kw}]'
    # 标志位（英文字段名兜底）
    if f.startswith('is_') or f.endswith('_flag'):
        return True, '字段名标志位模式'
    if f in ('delete_flag', 'del_flag', 'flag', 'status', 'deleteflag', 'deleted'):
        return True, '字段名标志位模式'
    # 统计数
    for kw in ('数量', '总数', '次数', '统计'):
        if kw in d:
            return True, f'命中[{kw}]'
    if f.endswith('_count') or f.endswith('_num') or f.startswith('count'):
        return True, '字段名统计数模式'
    return False, ''


def downgrade_check(fname, desc, cat, mark):
    """敏感命中后的降级判定：真实语义是日期/类型/id/标志位时降级，避免规则误判。"""
    d = (desc or '')
    f = (fname or '').lower()
    # "是否"开头是标志位（如"是否保留车牌"），非敏感值本身 → 一级
    if '是否' in d:
        return '非敏感数据', '不敏感', '一级', '语义为标志位(是否)，降级'
    # 证件类但注释是"有效期/期限" → 是日期非证件号 → 二级
    if mark == '四级' and ('有效期' in d or '期限' in d):
        return '业务数据', '低敏感', '二级', '语义为有效期/日期，降级'
    # "类型/车型/类别"且无"号/码" → 是分类值非敏感值 → 二级
    if ('类型' in d or '车型' in d or '类别' in d) and '号' not in d and '码' not in d:
        return '业务数据', '低敏感', '二级', '语义为类型/类别，降级'
    # 注释含 id 且字段名以 _id 结尾 → 外键 id → 二级
    if 'id' in d.lower() and f.endswith('_id'):
        return '业务数据', '低敏感', '二级', '语义为外键id，降级'
    return None


def classify(desc, tdesc, fname):
    full = (desc or '') + ' ' + (tdesc or '')
    f = (fname or '').lower()

    def finalize(cat, lvl, mark, why):
        dg = downgrade_check(fname, desc, cat, mark)
        return dg if dg else (cat, lvl, mark, why)

    for kw, cat in HIGH_CN.items():
        if kw in full:
            return finalize(cat, '高敏感', '四级', f'命中[{kw}]')
    for kw in HIGH_EN:
        if kw in f:
            cat = '身份鉴别信息' if kw in ('password', 'passwd', 'pwd') else '个人身份信息'
            return finalize(cat, '高敏感', '四级', f'字段名命中[{kw}]')
    for kw, cat in SENS_CN.items():
        if kw in full:
            if kw in ('住址',) and is_url_field(fname, full):
                continue
            return finalize(cat, '敏感', '三级', f'命中[{kw}]')
    for kw in SENS_EN:
        if kw in f:
            if kw in ('plate', 'carno'):
                cat = '车辆标识信息'
            elif kw == 'openid':
                cat = '个人身份信息'
            elif kw in ('birthday', 'contact'):
                cat = '个人基本资料'
            else:
                cat = '个人通信信息'
            return finalize(cat, '敏感', '三级', f'字段名命中[{kw}]')
    ok, why = is_level1(fname, desc)
    if ok:
        return '非敏感数据', '不敏感', '一级', why
    return '业务数据', '低敏感', '二级', ''


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else '数据资产清单.xlsx'
    prefix = sys.argv[2] if len(sys.argv) > 2 else ''

    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(x) if x is not None else '' for x in rows[0]]
    idx = locate_columns(header)

    def g(r, name):
        i = idx.get(name)
        return r[i] if i is not None else None

    hdr = ['ip', 'port', '实例名', '数据库名', '模式名', '数据表名', '字段名称',
           '字段描述', '表描述', '数据类别(修正)', '敏感级别', '分级标识',
           '判定依据', '建议管控措施']

    full_out = openpyxl.Workbook()
    fws = full_out.active
    fws.title = '分类分级清单'
    fws.append(hdr)

    sens_out = openpyxl.Workbook()
    sws = sens_out.active
    sws.title = '敏感字段待确认'
    sws.append(hdr)

    cnt, lv = Counter(), Counter()
    sens_cnt = 0

    for r in rows[1:]:
        if all((r[i] is None or str(r[i]).strip() == '') for i in range(len(r))):
            continue  # 空行跳过
        cat, lvl, mark, why = classify(g(r, '字段描述'), g(r, '表描述'), g(r, '字段名称'))
        cnt[cat] += 1
        lv[mark] += 1
        line = [g(r, c) for c in OPTIONAL] + [g(r, '数据表名'), g(r, '字段名称'),
                g(r, '字段描述'), g(r, '表描述'), cat, lvl, mark, why,
                MEASURE.get(cat, '')]
        fws.append(line)
        if mark in ('三级', '四级'):
            sens_cnt += 1
            sws.append(line)

    # 样式：越敏感越红 —— 一级白 / 二级浅黄 / 三级浅红 / 四级深红
    hfill = PatternFill('solid', fgColor='1F4E79')
    lv2 = PatternFill('solid', fgColor='FFF2CC')   # 二级 浅黄
    lv3 = PatternFill('solid', fgColor='FCE4E4')   # 三级 浅红
    lv4 = PatternFill('solid', fgColor='F4B7B7')   # 四级 深红
    fill_map = {'一级': None, '二级': lv2, '三级': lv3, '四级': lv4}
    widths = [14, 8, 10, 12, 10, 22, 24, 22, 24, 16, 10, 10, 16, 34]

    for ows in (fws, sws):
        for c in ows[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = hfill
            c.alignment = Alignment(horizontal='center')
        for row in ows.iter_rows(min_row=2):
            mark = row[11].value
            fill = fill_map.get(mark)
            if fill:
                for c in row:
                    c.fill = fill
        for i, w in enumerate(widths, 1):
            ows.column_dimensions[get_column_letter(i)].width = w
        ows.freeze_panes = 'A2'

    f1 = f'{prefix}分类分级清单-修正版.xlsx'
    f2 = f'{prefix}敏感字段清单-待人工确认.xlsx'
    full_out.save(f1)
    sens_out.save(f2)

    print(f'输入: {src}')
    print(f'分类分布: {dict(cnt)}')
    print(f'分级分布: {dict(lv)}')
    print(f'敏感字段(三级+四级): {sens_cnt} 条')
    print(f'已生成: {f1} / {f2}')


if __name__ == '__main__':
    main()
